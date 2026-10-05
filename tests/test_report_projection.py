import json
from pathlib import Path

import pytest
from docx import Document
from fastapi.testclient import TestClient

from portrait_eval.config import Settings
from portrait_eval.database import Database
from portrait_eval.product_api import create_app
from portrait_eval.report_projection import resolve_report_payload
from portrait_eval.reporting import ReportPayload, render_report_bundle
from portrait_eval.repository import Repository
from portrait_eval.workflow import finalize_quick_report


def fixture_payload() -> ReportPayload:
    a = {
        "claim_id": "a",
        "device_id": "device-a",
        "statement": "removed-secret",
        "claim_type": "strategy",
        "grade": "B",
        "evidence": ["G001"],
    }
    b = {**a, "claim_id": "b", "device_id": "device-b"}
    return ReportPayload(
        project_name="Projection",
        devices=["A", "B"],
        findings=[a, b],
        scene_results=[
            {
                "group_id": "G001",
                "findings": [a, b],
                "primary": {"raw": {"adapter": "heuristic"}, "observations": [a, b]},
            }
        ],
        device_profiles=[
            {"device_id": "device-a", "summary": "removed-secret", "claim_ids": ["a"]},
            {"device_id": "device-b", "summary": "removed-secret", "claim_ids": ["b"]},
        ],
        attributions=[a, b],
    )


def docx_text(path: Path) -> str:
    document = Document(path)
    return "\n".join(
        [p.text for p in document.paragraphs]
        + [c.text for t in document.tables for r in t.rows for c in r.cells]
    )


def test_ids_keep_other_device_and_remove_all_claim_representations(tmp_path: Path):
    projected = resolve_report_payload(
        fixture_payload(), [{"status": "rejected", "payload": {"claim_id": "a"}}]
    )
    assert [item["claim_id"] for item in projected.findings] == ["b"]
    assert [item["claim_id"] for item in projected.scene_results[0]["primary"]["observations"]] == [
        "b"
    ]
    assert [item["claim_id"] for item in projected.attributions] == ["b"]
    assert [item["device_id"] for item in projected.device_profiles] == ["device-b"]


def test_nested_legacy_reject_is_scene_and_device_scoped():
    payload = fixture_payload()
    for item in payload.findings:
        item.pop("claim_id")
    review = {
        "status": "insufficient_evidence",
        "payload": {
            "group_id": "G001",
            "primary_observation": {"device_id": "device-a", "statement": "removed-secret"},
        },
    }
    projected = resolve_report_payload(payload, [review])
    assert [item["device_id"] for item in projected.findings] == ["device-b"]
    assert [
        item["device_id"] for item in projected.scene_results[0]["primary"]["observations"]
    ] == ["device-b"]


def setup_report(tmp_path: Path, payload: ReportPayload):
    settings = Settings(
        database_url=f"sqlite:///{tmp_path / 'report.db'}", workspace=tmp_path / "workspace"
    )
    settings.prepare()
    client = TestClient(create_app(settings))
    project_id = client.post("/api/projects", json={"name": "Publication"}).json()["id"]
    database = Database(settings.database_url)
    bundle = render_report_bundle(
        payload, settings.workspace / "projects" / project_id / "reports", "0.1", "draft"
    )
    with database.session_factory() as session:
        repo = Repository(session)
        repo.save_report(project_id, "0.1", "draft", str(bundle["html"]))
        review = repo.create_review_item(project_id, "model_conflict", {"claim_id": "a"})
        repo.resolve_review_item(review.id, "rejected", "Unsupported")
    return client, database, project_id, bundle


def test_actual_api_finalize_json_html_docx_use_resolved_projection(tmp_path: Path):
    payload = fixture_payload()
    # Keep a distinct accepted statement, so no rejected text can hide in any renderer.
    for section in (
        payload.findings,
        payload.attributions,
        payload.scene_results[0]["findings"],
        payload.scene_results[0]["primary"]["observations"],
    ):
        section[1]["statement"] = "accepted-b"
    payload.device_profiles[1]["summary"] = "accepted-b"
    client, _database, project_id, _ = setup_report(tmp_path, payload)
    response = client.post(f"/api/projects/{project_id}/reports/finalize")
    assert response.status_code == 200, response.text
    final = Path(response.json()["html_path"])
    for content in (
        final.read_text(),
        final.with_suffix(".json").read_text(),
        docx_text(final.with_suffix(".docx")),
    ):
        assert "removed-secret" not in content
        assert "accepted-b" in content
        assert "Heuristic adapter evidence is provisional" in content
    assert json.loads(final.with_suffix(".json").read_text())["review_summary"]["open"] == 0


def test_missing_json_fails_safely(tmp_path: Path):
    client, _, project_id, draft = setup_report(tmp_path, fixture_payload())
    draft["json"].unlink()
    assert client.post(f"/api/projects/{project_id}/reports/finalize").status_code == 409
    assert not (draft["html"].parent / "final-v1.0.html").exists()


def test_quick_applies_rejections_but_discloses_open_reviews(tmp_path: Path):
    _client, database, project_id, _ = setup_report(tmp_path, fixture_payload())
    with database.session_factory() as session:
        repo = Repository(session)
        repo.create_review_item(project_id, "open_question", {"statement": "pending"})
        final = finalize_quick_report(repo, project_id)
    path = Path(final["html_path"])
    data = json.loads(path.with_suffix(".json").read_text())
    assert data["review_summary"]["open"] == 1
    assert data["review_summary"]["skipped_mandatory_gates"] is True
    assert [item["claim_id"] for item in data["findings"]] == ["b"]
    for content in (path.read_text(), docx_text(path.with_suffix(".docx"))):
        assert "1 open reviews" in content
        assert "mandatory review gates skipped" in content
        assert "review gates resolved" not in content


def test_ambiguous_text_only_legacy_decision_fails_closed():

    with pytest.raises(ValueError, match="ambiguous device"):
        resolve_report_payload(
            fixture_payload(), [{"status": "rejected", "payload": {"statement": "removed-secret"}}]
        )


def test_unique_text_only_legacy_decision_is_resolved():
    payload = ReportPayload(
        project_name="legacy", devices=["A"], findings=[{"device_id": "a", "statement": "unique"}]
    )
    assert not resolve_report_payload(
        payload, [{"status": "rejected", "payload": {"statement": "unique"}}]
    ).findings


def test_synthetic_provenance_true_and_false_are_distinguished(tmp_path: Path):
    from portrait_eval.report_projection import publication_notes

    payload = ReportPayload(
        project_name="provenance", devices=[], findings=[], scene_results=[{"synthetic": False}]
    )
    assert not any(note.startswith("Synthetic evidence") for note in publication_notes(payload))
    payload.scene_results[0]["synthetic"] = True
    bundle = render_report_bundle(payload, tmp_path, "0.1", "draft")
    for content in (
        bundle["html"].read_text(),
        bundle["json"].read_text(),
        docx_text(bundle["docx"]),
    ):
        assert "Synthetic evidence is demonstration-only" in content


@pytest.mark.parametrize("mode", ["professional", "quick"])
def test_actual_api_excludes_provider_bodies_but_retains_immutable_audit_and_hashes(
    tmp_path: Path, mode: str
):
    import hashlib

    payload = fixture_payload()
    for section in (
        payload.findings,
        payload.attributions,
        payload.scene_results[0]["findings"],
        payload.scene_results[0]["primary"]["observations"],
    ):
        section[1]["statement"] = "accepted-b"
    payload.device_profiles[1]["summary"] = "accepted-b"
    provider_body = json.dumps(
        {"observations": [{"device_id": "device-a", "statement": "removed-secret"}]}
    )
    raw = {
        "adapter": "heuristic",
        "model": "model-v1",
        "choices": [{"message": {"content": provider_body}}],
        "input_trace": {"prompt_sha256": "prompt-hash", "ordered_image_sha256": ["image-hash"]},
        "prompt": "removed-secret prompt",
        "usage": {"total_tokens": 30},
    }
    payload.scene_results[0]["primary"]["raw"] = raw
    client, database, project_id, draft = setup_report(tmp_path, payload)
    original = draft["json"].read_bytes()
    with database.session_factory() as session:
        Repository(session).save_analysis(project_id, "primary_vlm", {"raw": raw})
    if mode == "professional":
        response = client.post(f"/api/projects/{project_id}/reports/finalize")
        assert response.status_code == 200, response.text
        path = Path(response.json()["html_path"])
    else:
        with database.session_factory() as session:
            path = Path(finalize_quick_report(Repository(session), project_id)["html_path"])
    for content in (
        path.read_text(),
        path.with_suffix(".json").read_text(),
        docx_text(path.with_suffix(".docx")),
    ):
        assert "removed-secret" not in content
        assert "accepted-b" in content
        assert "Heuristic adapter evidence is provisional" in content
    published = json.loads(path.with_suffix(".json").read_text())["scene_results"][0]["primary"][
        "raw"
    ]
    assert published["adapter"] == "heuristic" and published["model"] == "model-v1"
    assert published["input_trace"]["prompt_sha256"] == "prompt-hash"
    assert published["input_trace"]["ordered_image_sha256"] == ["image-hash"]
    assert (
        published["response_sha256"]
        == hashlib.sha256(
            json.dumps(raw, sort_keys=True, separators=(",", ":"), ensure_ascii=False).encode()
        ).hexdigest()
    )
    assert "choices" not in published and "prompt" not in published
    assert draft["json"].read_bytes() == original
    with database.session_factory() as session:
        assert (
            Repository(session).list_analysis(project_id, "primary_vlm")[0]["payload"]["raw"] == raw
        )


def test_authoritative_claim_id_preserves_distinct_id_and_filters_idless_equivalent():
    a = {
        "claim_id": "a",
        "device_id": "d",
        "dimension": "tone",
        "statement": "same",
        "evidence": ["G001"],
    }
    b = {**a, "claim_id": "b", "dimension": "color"}
    historical = {key: value for key, value in a.items() if key != "claim_id"}
    payload = ReportPayload(
        project_name="ids",
        devices=["D"],
        findings=[a, b],
        scene_results=[{"group_id": "G001", "primary": {"observations": [a, b, historical]}}],
    )
    review = {
        "status": "rejected",
        "payload": {"claim_id": "a", "device_id": "d", "statement": "same"},
    }
    projected = resolve_report_payload(payload, [review])
    assert [item["claim_id"] for item in projected.findings] == ["b"]
    assert projected.scene_results[0]["primary"]["observations"] == [b]


def test_legacy_ambiguity_respects_scene_and_dimension_scope():
    a = {"device_id": "a", "dimension": "tone", "statement": "same", "evidence": ["G001"]}
    b = {**a, "device_id": "b", "evidence": ["G002"]}
    c = {**a, "device_id": "c", "dimension": "color"}
    payload = ReportPayload(
        project_name="scopes",
        devices=["A", "B", "C"],
        findings=[a, b, c],
        scene_results=[
            {"group_id": "G001", "primary": {"observations": [a, c]}},
            {"group_id": "G002", "primary": {"observations": [b]}},
        ],
    )
    review = {
        "status": "rejected",
        "payload": {
            "group_id": "G001",
            "primary_observation": {"statement": "same", "dimension": "tone"},
        },
    }
    projected = resolve_report_payload(payload, [review])
    assert projected.findings == [b, c]
    assert projected.scene_results[0]["primary"]["observations"] == [c]
    assert projected.scene_results[1]["primary"]["observations"] == [b]
