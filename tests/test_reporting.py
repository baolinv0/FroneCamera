from pathlib import Path

import pytest
from docx import Document
from PIL import Image

from portrait_eval.docx_reporting import PACKAGE_TEMPLATE
from portrait_eval.reporting import ReportPayload, render_html_report


def test_report_contains_traceable_claims(tmp_path: Path) -> None:
    payload = ReportPayload(
        project_name="Demo",
        devices=["A", "B"],
        findings=[{"statement": "A is brighter in G001", "evidence": ["G001"]}],
        limitations=["Single capture"],
    )
    output = tmp_path / "report.html"
    render_html_report(payload, output)
    text = output.read_text(encoding="utf-8")
    assert "A is brighter in G001" in text
    assert "G001" in text
    assert "Single capture" in text


def test_final_report_bundle_has_final_banner_and_preserves_draft(tmp_path: Path) -> None:
    from portrait_eval.reporting import render_report_bundle

    payload = ReportPayload(project_name="Demo", devices=["A", "B"], findings=[])
    draft = render_report_bundle(payload, tmp_path, version="0.1", status="draft")
    final = render_report_bundle(payload, tmp_path, version="1.0", status="final")
    assert "DRAFT" in draft["html"].read_text(encoding="utf-8")
    final_text = final["html"].read_text(encoding="utf-8")
    assert "FINAL REPORT" in final_text
    assert "DRAFT" not in final_text
    assert draft["html"].exists()
    assert draft["docx"].exists()
    assert final["docx"].exists()


def test_report_bundle_generates_reference_derived_docx(tmp_path: Path) -> None:
    from portrait_eval.reporting import render_report_bundle

    image_path = tmp_path / "source-image.jpg"
    Image.new("RGB", (640, 480), "#f97316").save(image_path)
    payload = ReportPayload(
        project_name="六机前摄评测",
        devices=["Device A", "Device B"],
        findings=[
            {
                "device_id": "device-a",
                "statement": "Device A preserves highlight structure in G001.",
                "grade": "A",
                "status": "CONFIRMED",
                "evidence": ["G001"],
            }
        ],
        scene_results=[
            {
                "group_id": "G001",
                "findings": [
                    {
                        "device_id": "device-a",
                        "statement": "Highlight structure remains visible.",
                        "grade": "A",
                    }
                ],
            }
        ],
        visual_assets=[
            {
                "scene_id": "G001",
                "device_id": "device-a",
                "device_name": "Device A",
                "path": str(image_path),
            }
        ],
        device_profiles=[
            {
                "device_id": "device-a",
                "device_name": "Device A",
                "label": "Highlight-Safe Realism",
                "summary": "Protects light-source structure.",
            },
            {
                "device_id": "device-b",
                "device_name": "Device B",
                "label": "Subject Priority",
                "summary": "Raises face readability.",
            },
        ],
        limitations=["JPEG output does not uniquely identify internal ISP mechanisms."],
    )

    bundle = render_report_bundle(payload, tmp_path / "reports", "0.1", "draft")
    generated = Document(bundle["docx"])
    text = "\n".join(
        [paragraph.text for paragraph in generated.paragraphs]
        + [cell.text for table in generated.tables for row in table.rows for cell in row.cells]
    )

    assert "六机前摄评测 前摄影调评测" in text
    assert "Executive Summary" in text
    assert "Scene 1" in text
    assert "Final Verdict" in text
    assert "Appendix" in text
    assert len(generated.inline_shapes) >= 2
    assert any(
        paragraph.style.name == "Heading 1"
        for table in generated.tables
        for row in table.rows
        for cell in row.cells
        for paragraph in cell.paragraphs
    )
    drawing_properties = generated._element.xpath(".//wp:docPr")
    assert drawing_properties
    assert all(item.get("descr") for item in drawing_properties)
    section = generated.sections[0]
    assert section.page_width.inches == pytest.approx(8.2681, abs=0.01)
    assert section.left_margin.inches == pytest.approx(0.5313, abs=0.01)
    assert "PAGE" in section.footer._element.xml

    serialized = bundle["json"].read_text(encoding="utf-8")
    assert str(image_path) not in serialized
    assert "assets/001-G001-Device-A.jpg" in serialized
    assert (bundle["html"].parent / "assets" / "001-G001-Device-A.jpg").exists()


def test_packaged_template_contains_no_reference_content_or_media() -> None:
    import zipfile

    template = Document(PACKAGE_TEMPLATE)
    assert not template.paragraphs
    assert not template.tables
    assert not template.inline_shapes
    with zipfile.ZipFile(PACKAGE_TEMPLATE) as archive:
        assert not any(name.startswith("word/media/") for name in archive.namelist())
        assert not any(name.startswith("customXml/") for name in archive.namelist())
        assert "docProps/thumbnail.jpeg" not in archive.namelist()
        document_xml = archive.read("word/document.xml").decode("utf-8")
        assert "六款旗舰手机" not in document_xml
