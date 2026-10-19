import json
import re

import pytest
from docx import Document
from PIL import Image
from pydantic import ValidationError

from portrait_eval.reporting import ReportPayload, render_report_bundle


def report_fixture(tmp_path):
    image = tmp_path / "sample.png"
    Image.new("RGB", (120, 160), (80, 110, 140)).save(image)
    claim = {
        "claim_id": "claim:one",
        "claim_version": "0.2:test",
        "device_id": "a",
        "statement": "Visible luma difference, limited to room-1",
        "grade": "C",
        "status": "DISPUTED",
        "attribute": "encoded_luma_mean",
        "evidence": ["G001"],
        "supporting_scene_ids": ["G001"],
        "contradicting_scene_ids": ["G002"],
        "evidence_counts": {
            "support": 1,
            "counterexample": 1,
            "unknown": 0,
            "incomparable": 0,
            "total": 2,
        },
    }
    payload = ReportPayload(
        project_name="Evidence report",
        devices=["A", "B"],
        schema_version="0.2",
        findings=[],
        all_claims=[claim],
        scene_results=[
            {
                "group_id": "G001",
                "scene_family": "room-1",
                "audit": {"challenge": {"primary": "mixed_light"}},
                "findings": [claim],
            },
            {
                "group_id": "G002",
                "audit": {"challenge": {"primary": "high_contrast"}},
                "findings": [],
            },
        ],
        visual_assets=[{"scene_id": "G001", "device_id": "a", "path": str(image)}],
        attributions=[
            {
                "statement": "SPECIFIC OBSERVABLE EFFECT",
                "candidate_causes": ["CANDIDATE FOCUS"],
                "alternative_explanations": ["motion"],
                "primary_layer": "indeterminate",
                "minimal_experiment": "FIXED DISTANCE RESHOOT",
            }
        ],
        capture_bias=[
            {
                "status": "POSSIBLE_CAPTURE_BIAS",
                "reasons": ["SPECIFIC BIAS REASON"],
                "recommended_actions": ["ALTERNATE DEVICES"],
            }
        ],
    )
    return payload


def document_text(path):
    d = Document(path)
    return "\n".join(
        [p.text for p in d.paragraphs] + [c.text for t in d.tables for r in t.rows for c in r.cells]
    )


def test_html_docx_share_nonempty_evidence_and_actual_module_fields(tmp_path):
    bundle = render_report_bundle(report_fixture(tmp_path), tmp_path / "reports", "0.2", "final")
    h = bundle["html"].read_text()
    d = document_text(bundle["docx"])
    for token in [
        "SPECIFIC OBSERVABLE EFFECT",
        "CANDIDATE FOCUS",
        "SPECIFIC BIAS REASON",
        "ALTERNATE DEVICES",
        "Visible luma difference",
        "G002",
    ]:
        assert token in h
        assert token in d
    assert "模型辅助观察报告" in h and "模型辅助观察报告" in d
    assert "review gates resolved" not in h and "review gates resolved" not in d
    assert "<img " in h
    ids = set(re.findall(r'id="([^"]+)"', h))
    assert all(target in ids for target in re.findall(r'href="#([^"]+)"', h))
    for asset in re.findall(r'<img[^>]+src="([^"]+)"', h):
        assert (bundle["html"].parent / asset).is_file()
    data = json.loads(bundle["json"].read_text())
    assert len(data["all_claims"]) == 1
    assert data["all_claims"][0]["claim_id"] == "claim:one"


def test_new_report_contract_rejects_orphan_claims():
    with pytest.raises(ValidationError, match="claim_id"):
        ReportPayload(
            project_name="bad", devices=[], schema_version="0.2", findings=[{"statement": "orphan"}]
        )


def test_external_evidence_join_requires_device_claim_and_version():
    from portrait_eval.evidence import related_external

    a = {
        "claim_id": "a",
        "claim_version": "v1",
        "device_id": "D1",
        "attribute": "luma",
        "statement": "same",
    }
    records = [
        {**a, "url": "https://one.test"},
        {**a, "claim_id": "b", "device_id": "D2"},
        {**a, "claim_version": "v2"},
    ]
    assert len(related_external(a, records)) == 1


def test_review_resolution_is_claim_bound_and_updates_scene_narrative(tmp_path):
    from portrait_eval.report_model import apply_review_decisions

    payload = report_fixture(tmp_path)
    a = {
        **payload.all_claims[0],
        "grade": "B",
        "status": "SUPPORTED_SCOPED",
        "contradicting_scene_ids": [],
    }
    b = {**a, "claim_id": "claim:two", "device_id": "b"}
    payload.findings = [a, b]
    payload.all_claims = [a, b]
    payload.scene_results[0]["findings"] = [a, b]
    result = apply_review_decisions(
        payload, [{"status": "rejected", "category": "low_confidence_claim", "payload": a}]
    )
    assert [c["claim_id"] for c in result.findings] == ["claim:two"]
    assert result.scene_results[0]["findings"][0]["status"] == "HUMAN_REJECTED"
    assert payload.scene_results[0]["findings"][0]["status"] == "SUPPORTED_SCOPED"


def test_served_report_images_use_guarded_asset_routes(tmp_path):
    from fastapi.testclient import TestClient

    from portrait_eval.config import Settings
    from portrait_eval.product_api import create_app
    from portrait_eval.repository import Repository

    settings = Settings(
        database_url=f"sqlite:///{tmp_path / 'web.db'}",
        workspace=tmp_path / "workspace",
        api_token="secret",
        report_share_secret="test-secret",
    )
    app = create_app(settings)
    client = TestClient(app)
    with app.state.database.session_factory() as session:
        repo = Repository(session)
        project = repo.create_project("web report")
        bundle = render_report_bundle(
            report_fixture(tmp_path), settings.workspace / "reports" / "run1", "0.2", "draft"
        )
        row = repo.save_report(project.id, "0.2", "draft", str(bundle["html"]))
        report_id = row.id
    headers = {"Authorization": "Bearer secret"}
    private = client.get(f"/api/reports/{report_id}/html", headers=headers)
    src = re.findall(r'<img[^>]+src="([^"]+)"', private.text)[0]
    assert src.startswith(f"/api/reports/{report_id}/assets/")
    assert client.get(src, headers=headers).status_code == 200
    assert client.get(src).status_code == 401
    share = client.post(f"/api/reports/{report_id}/share", headers=headers).json()
    public = client.get(share["url"])
    src = re.findall(r'<img[^>]+src="([^"]+)"', public.text)[0].replace("&amp;", "&")
    assert src.startswith(f"/reports/{report_id}/assets/")
    assert client.get(src).status_code == 200
    assert client.get(src.split("?")[0] + "?token=bad").status_code == 403
    outside = tmp_path / "outside.png"
    outside.write_bytes(b"private bytes")
    (bundle["html"].parent / "assets" / "escape.png").symlink_to(outside)
    token = share["token"]
    assert client.get(f"/reports/{report_id}/assets/escape.png?token={token}").status_code == 403


def test_rejecting_external_source_keeps_internal_visual_claim(tmp_path):
    from portrait_eval.report_model import apply_review_decisions

    payload = report_fixture(tmp_path)
    claim = {
        **payload.all_claims[0],
        "grade": "B",
        "status": "SUPPORTED_SCOPED",
        "contradicting_scene_ids": [],
    }
    payload.findings = [claim]
    payload.all_claims = [claim]
    external = {**claim, "url": "https://source.test", "evidence_scope": "background"}
    payload.external_validation = [external]
    reviewed = apply_review_decisions(
        payload,
        [{"status": "rejected", "category": "external_corroboration_review", "payload": external}],
    )
    assert len(reviewed.findings) == 1
    assert reviewed.all_claims[0]["status"] == "SUPPORTED_SCOPED"
    assert not reviewed.external_validation


def test_docx_embeds_every_native_face_and_keeps_full_frame(tmp_path):
    import numpy as np

    from portrait_eval.docx_reporting import _create_montage

    payload = report_fixture(tmp_path)
    image = Image.new("RGB", (100, 400), "green")
    from PIL import ImageDraw

    draw = ImageDraw.Draw(image)
    draw.rectangle((0, 0, 99, 50), fill="red")
    draw.rectangle((0, 350, 99, 399), fill="blue")
    portrait = tmp_path / "portrait.png"
    image.save(portrait)
    montage, _ = _create_montage([{"path": str(portrait)}], tmp_path, tmp_path / "montage.jpg")
    pixels = np.asarray(Image.open(montage))
    assert ((pixels[:, :, 0] > 220) & (pixels[:, :, 1] < 30)).any()
    assert ((pixels[:, :, 2] > 220) & (pixels[:, :, 1] < 30)).any()
    payload.scene_results[0]["metrics"] = {
        "a": {
            "faces": [
                {
                    "face_id": f"F00{i}",
                    "role": "SECONDARY",
                    "bbox_xywh": [0, 0, 100, 400],
                    "native_crop_path": str(portrait),
                }
                for i in (1, 2)
            ]
        }
    }
    bundle = render_report_bundle(payload, tmp_path / "face-reports", "0.2", "draft")
    doc = Document(bundle["docx"])
    text = document_text(bundle["docx"])
    assert "F001" in text and "F002" in text
    assert len(doc.inline_shapes) >= 4


def test_reference_footer_does_not_leak_fixed_device_count(tmp_path):
    bundle = render_report_bundle(report_fixture(tmp_path), tmp_path / "footer", "0.2", "draft")
    doc = Document(bundle["docx"])
    text = "\n".join(p.text for section in doc.sections for p in section.footer.paragraphs)
    text += "\n".join(
        c.text
        for section in doc.sections
        for table in section.footer.tables
        for row in table.rows
        for c in row.cells
    )
    assert "六款旗舰手机" not in text
    assert "PAGE" in doc.sections[0].footer._element.xml
