from __future__ import annotations

import math
import tempfile
from collections import Counter, defaultdict
from collections.abc import Iterable
from pathlib import Path
from typing import TYPE_CHECKING, Any

from docx import Document
from docx.document import Document as DocxDocument
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont, ImageOps

if TYPE_CHECKING:
    from portrait_eval.reporting import ReportPayload


INK = "202124"
ACCENT = "F97316"
WHITE = "FFFFFF"
MUTED = "666666"
RULE = "D6D6D6"
CARD = "F7F7F7"
WARM = "FFF3E8"
FONT = "Noto Sans CJK SC"
CONTENT_WIDTH = 7.20
PACKAGE_TEMPLATE = Path(__file__).parent / "assets" / "report_reference_template.docx"


def _rgb(hex_value: str) -> RGBColor:
    return RGBColor.from_string(hex_value)


def _set_run_font(
    run: Any,
    *,
    size: float,
    color: str = INK,
    bold: bool = False,
    italic: bool = False,
) -> None:
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.color.rgb = _rgb(color)
    run.font.bold = bold
    run.font.italic = italic
    properties = run._element.get_or_add_rPr()
    fonts = properties.get_or_add_rFonts()
    fonts.set(qn("w:ascii"), FONT)
    fonts.set(qn("w:hAnsi"), FONT)
    fonts.set(qn("w:eastAsia"), FONT)


def _ensure_style(
    document: DocxDocument,
    name: str,
    *,
    size: float,
    color: str = INK,
    bold: bool = False,
    space_before: float = 0,
    space_after: float = 4,
    line_spacing: float = 1.08,
) -> Any:
    styles = document.styles
    style = styles[name] if name in styles else styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    style.font.name = FONT
    style.font.size = Pt(size)
    style.font.color.rgb = _rgb(color)
    style.font.bold = bold
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:ascii"), FONT)
    fonts.set(qn("w:hAnsi"), FONT)
    fonts.set(qn("w:eastAsia"), FONT)
    paragraph = style.paragraph_format
    paragraph.space_before = Pt(space_before)
    paragraph.space_after = Pt(space_after)
    paragraph.line_spacing = line_spacing
    return style


def _configure_styles(document: DocxDocument) -> None:
    _ensure_style(document, "Normal", size=8.5, space_after=3, line_spacing=1.08)
    _ensure_style(
        document,
        "Report Eyebrow",
        size=9,
        color=ACCENT,
        bold=True,
        space_after=5,
    )
    _ensure_style(document, "Report Title", size=29, bold=True, space_after=6, line_spacing=0.95)
    _ensure_style(document, "Report Subtitle", size=11.5, color=MUTED, space_after=14)
    _ensure_style(document, "Report Body", size=8.5, space_after=4)
    _ensure_style(document, "Report Small", size=7, color=MUTED, space_after=2)
    _ensure_style(document, "Report Caption", size=7, color=MUTED, space_before=2, space_after=8)
    _ensure_style(document, "Report H1", size=20, bold=True, space_before=8, space_after=5)
    _ensure_style(document, "Report H2", size=12.5, bold=True, space_before=6, space_after=3)
    _ensure_style(document, "Report Label", size=8, color=ACCENT, bold=True, space_after=2)
    for standard_name, report_name in (
        ("Title", "Report Title"),
        ("Subtitle", "Report Subtitle"),
        ("Heading 1", "Report H1"),
        ("Heading 2", "Report H2"),
    ):
        source = document.styles[report_name]
        target = document.styles[standard_name]
        target.font.name = source.font.name
        target.font.size = source.font.size
        target.font.color.rgb = source.font.color.rgb
        target.font.bold = source.font.bold
        target.paragraph_format.space_before = source.paragraph_format.space_before
        target.paragraph_format.space_after = source.paragraph_format.space_after
        target.paragraph_format.line_spacing = source.paragraph_format.line_spacing
    if "List Bullet" in document.styles:
        bullet = document.styles["List Bullet"]
        bullet.font.name = FONT
        bullet.font.size = Pt(8)
        bullet.font.color.rgb = _rgb(INK)
        bullet.paragraph_format.left_indent = Inches(0.18)
        bullet.paragraph_format.first_line_indent = Inches(-0.12)
        bullet.paragraph_format.space_after = Pt(2)


def _shade_cell(cell: Any, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)
    shading.set(qn("w:val"), "clear")


def _set_cell_margins(
    cell: Any, *, top: int = 110, start: int = 120, bottom: int = 110, end: int = 120
) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_cell_border(cell: Any, *, color: str = RULE, size: int = 6) -> None:
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)
    for edge in ("top", "start", "bottom", "end", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        node = borders.find(tag)
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), str(size))
        node.set(qn("w:color"), color)


def _set_repeat_table_header(row: Any) -> None:
    properties = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    properties.append(header)


def _set_table_geometry(table: Any, widths: Iterable[float]) -> None:
    requested_widths = list(widths)
    target_twips = round(CONTENT_WIDTH * 1440) - 120
    requested_total = sum(requested_widths)
    width_twips = [round(width / requested_total * target_twips) for width in requested_widths[:-1]]
    width_twips.append(target_twips - sum(width_twips))
    widths_list = [width / 1440 for width in width_twips]
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    properties = table._tbl.tblPr
    layout = properties.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        properties.append(layout)
    layout.set(qn("w:type"), "fixed")
    total = target_twips
    width_node = properties.first_child_found_in("w:tblW")
    if width_node is None:
        width_node = OxmlElement("w:tblW")
        properties.append(width_node)
    width_node.set(qn("w:w"), str(total))
    width_node.set(qn("w:type"), "dxa")
    indent = properties.first_child_found_in("w:tblInd")
    if indent is None:
        indent = OxmlElement("w:tblInd")
        properties.append(indent)
    indent.set(qn("w:w"), "120")
    indent.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width_twip in width_twips:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(width_twip))
        grid.append(column)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width_inches = widths_list[min(index, len(widths_list) - 1)]
            twips = width_twips[min(index, len(width_twips) - 1)]
            cell.width = Inches(width_inches)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell_properties = cell._tc.get_or_add_tcPr()
            tc_width = cell_properties.first_child_found_in("w:tcW")
            if tc_width is None:
                tc_width = OxmlElement("w:tcW")
                cell_properties.append(tc_width)
            tc_width.set(qn("w:w"), str(twips))
            tc_width.set(qn("w:type"), "dxa")
            _set_cell_margins(cell)


def _clear_cell(cell: Any) -> Any:
    paragraph = cell.paragraphs[0]
    paragraph.clear()
    for extra in cell.paragraphs[1:]:
        extra._element.getparent().remove(extra._element)
    return paragraph


def _cell_text(
    cell: Any,
    text: str,
    *,
    size: float = 8.5,
    color: str = INK,
    bold: bool = False,
    alignment: WD_ALIGN_PARAGRAPH = WD_ALIGN_PARAGRAPH.LEFT,
) -> None:
    paragraph = _clear_cell(cell)
    paragraph.alignment = alignment
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.05
    run = paragraph.add_run(text)
    _set_run_font(run, size=size, color=color, bold=bold)


def _add_spacer(document: DocxDocument, points: float = 5) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = Pt(points)


def _set_picture_alt(picture: Any, title: str, description: str) -> None:
    properties = picture._inline.docPr
    properties.set("title", title)
    properties.set("descr", description)


def _add_cover(document: DocxDocument, payload: ReportPayload, status: str) -> None:
    eyebrow = document.add_paragraph(style="Report Eyebrow")
    eyebrow.add_run("SAMPLE-BASED CAMERA TONE ASSESSMENT")
    title = document.add_paragraph(style="Report Title")
    title.add_run(payload.report_title or f"{payload.project_name} 前摄影调评测")
    subtitle = document.add_paragraph(style="Report Subtitle")
    subtitle.add_run(
        payload.report_subtitle or "证据优先叙事 × 场景样张比较 × 硬件 / 采集 / 重建 / 渲染归因边界"
    )

    stats = document.add_table(rows=1, cols=3)
    _set_table_geometry(stats, [2.4, 2.4, 2.4])
    values = (
        (str(len(payload.devices)), "款机型", " / ".join(payload.devices[:4])),
        (str(len(payload.visual_assets)), "张代表样张", "嵌入报告用于复核"),
        (str(len(payload.scene_results)), "类场景", "按实际匹配组生成"),
    )
    for cell, (value, label, detail) in zip(stats.rows[0].cells, values, strict=True):
        _shade_cell(cell, CARD)
        _set_cell_border(cell, color=WHITE, size=10)
        paragraph = _clear_cell(cell)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(value)
        _set_run_font(run, size=20, color=ACCENT, bold=True)
        paragraph = cell.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(label)
        _set_run_font(run, size=9, bold=True)
        paragraph = cell.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(detail)
        _set_run_font(run, size=7, color=MUTED)
    _add_spacer(document, 6)
    banner = (
        "FINAL REPORT — review gates resolved."
        if status == "final"
        else "DRAFT — verify review gates before external distribution."
    )
    _add_callout(
        document,
        "报告性质",
        f"{banner} 结果仅适用于本次提交样张、设备固件、拍摄模式与场景。"
        "评分与归因必须保持样张证据可追溯。",
        accent=status == "final",
    )


def _add_section_band(document: DocxDocument, number: str, title: str, strapline: str) -> None:
    _add_spacer(document, 7)
    table = document.add_table(rows=1, cols=2)
    _set_table_geometry(table, [0.62, 6.58])
    number_cell, title_cell = table.rows[0].cells
    _shade_cell(number_cell, ACCENT)
    _shade_cell(title_cell, INK)
    _set_cell_border(number_cell, color=ACCENT, size=0)
    _set_cell_border(title_cell, color=INK, size=0)
    _cell_text(
        number_cell,
        number,
        size=12.5,
        color=WHITE,
        bold=True,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
    )
    paragraph = _clear_cell(title_cell)
    paragraph.style = document.styles["Heading 1"]
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run(title)
    _set_run_font(run, size=20, color=WHITE, bold=True)
    strap = title_cell.add_paragraph()
    strap.paragraph_format.space_before = Pt(1)
    strap.paragraph_format.space_after = Pt(0)
    run = strap.add_run(strapline)
    _set_run_font(run, size=8, color="D6D6D6")
    _add_spacer(document, 4)


def _add_callout(
    document: DocxDocument,
    label: str,
    text: str,
    *,
    accent: bool = False,
) -> None:
    table = document.add_table(rows=1, cols=2)
    _set_table_geometry(table, [1.25, 5.95])
    label_cell, body_cell = table.rows[0].cells
    _shade_cell(label_cell, ACCENT if accent else INK)
    _shade_cell(body_cell, WARM if accent else CARD)
    _set_cell_border(label_cell, color=WHITE, size=6)
    _set_cell_border(body_cell, color=WHITE, size=6)
    _cell_text(
        label_cell,
        label,
        size=8,
        color=WHITE,
        bold=True,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
    )
    _cell_text(body_cell, text or "暂无可发布证据。", size=8.5)
    _add_spacer(document, 4)


def _safe_text(value: Any, default: str = "") -> str:
    if value is None:
        return default
    return str(value).strip()


def _approved_findings(payload: ReportPayload) -> list[dict[str, Any]]:
    return [
        item
        for item in payload.findings
        if _safe_text(item.get("grade"), "B") in {"A", "B"}
        and _safe_text(item.get("status"), "SUPPORTED")
        not in {"DISPUTED", "LOW_CONFIDENCE", "INSUFFICIENT_SCENE_COVERAGE"}
    ]


def _add_executive_summary(document: DocxDocument, payload: ReportPayload) -> None:
    _add_section_band(
        document,
        "00",
        "Executive Summary — 证据驱动的成像路线",
        "先给出被证据支持的产品倾向，再进入逐场景复核。",
    )
    findings = _approved_findings(payload)
    profiles = payload.device_profiles or [
        {"device_name": device, "label": "Evidence Profile", "summary": "等待更多跨场景证据。"}
        for device in payload.devices
    ]
    cards = document.add_table(rows=math.ceil(max(len(profiles), 1) / 2), cols=2)
    _set_table_geometry(cards, [3.6, 3.6])
    for index, cell in enumerate(cell for row in cards.rows for cell in row.cells):
        _shade_cell(cell, WHITE)
        _set_cell_border(cell, color=RULE)
        if index >= len(profiles):
            _cell_text(cell, "", size=8)
            continue
        profile = profiles[index]
        name = _safe_text(profile.get("device_name") or profile.get("name"), f"Device {index + 1}")
        label = _safe_text(profile.get("label") or profile.get("profile"), "Evidence Profile")
        summary = _safe_text(profile.get("summary"))
        if not summary:
            matching = [
                _safe_text(item.get("statement"))
                for item in findings
                if _safe_text(item.get("device_name") or item.get("device_id"))
                in {name, _safe_text(profile.get("device_id"))}
            ]
            summary = "；".join(matching[:2]) or "当前没有足够的跨场景结论。"
        paragraph = _clear_cell(cell)
        run = paragraph.add_run(name)
        _set_run_font(run, size=11.5, bold=True)
        paragraph = cell.add_paragraph()
        run = paragraph.add_run(label)
        _set_run_font(run, size=8, color=ACCENT, bold=True)
        paragraph = cell.add_paragraph()
        run = paragraph.add_run(summary)
        _set_run_font(run, size=8.5)
    _add_spacer(document, 5)
    summary_text = "；".join(_safe_text(item.get("statement")) for item in findings[:4])
    _add_callout(document, "执行摘要", summary_text or "尚无通过证据门槛的结论。", accent=True)


def _score_rows(payload: ReportPayload) -> tuple[list[dict[str, Any]], bool]:
    if payload.device_scores:
        rows = []
        for item in payload.device_scores:
            rows.append(
                {
                    "device_name": _safe_text(
                        item.get("device_name") or item.get("device_id"), "Device"
                    ),
                    "value": float(item.get("relative_score", item.get("score", 0))),
                    "confidence": item.get("confidence"),
                    "label": "相对评分",
                }
            )
        return rows, True

    profiles = {
        _safe_text(item.get("device_id")): _safe_text(item.get("device_name") or item.get("name"))
        for item in payload.device_profiles
    }
    counts: Counter[str] = Counter()
    for finding in _approved_findings(payload):
        identity = _safe_text(finding.get("device_name") or finding.get("device_id"))
        counts[profiles.get(identity, identity)] += 1
    rows = [
        {
            "device_name": device,
            "value": float(counts.get(device, 0)),
            "confidence": None,
            "label": "通过门槛的结论数",
        }
        for device in payload.devices
    ]
    return rows, False


def _font(size: int, *, bold: bool = False) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    candidates = (
        Path("C:/Windows/Fonts/msyhbd.ttc" if bold else "C:/Windows/Fonts/msyh.ttc"),
        Path("C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf"),
    )
    for candidate in candidates:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size=size)
    return ImageFont.load_default()


def _create_score_chart(rows: list[dict[str, Any]], output: Path, *, quality_score: bool) -> Path:
    width = 1800
    row_height = 105
    height = 130 + max(len(rows), 1) * row_height
    canvas = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(canvas)
    draw.text(
        (35, 24),
        "相对评分" if quality_score else "证据覆盖",
        fill=f"#{INK}",
        font=_font(38, bold=True),
    )
    draw.text(
        (35, 72),
        "批次内相对值" if quality_score else "通过 A/B 证据门槛的结论数量，不代表画质分数",
        fill=f"#{MUTED}",
        font=_font(24),
    )
    maximum = max((float(item["value"]) for item in rows), default=1.0) or 1.0
    if quality_score:
        maximum = max(maximum, 100.0)
    for index, item in enumerate(rows):
        y = 130 + index * row_height
        draw.text((35, y + 26), item["device_name"], fill=f"#{INK}", font=_font(27, bold=True))
        bar_left, bar_right = 520, 1660
        draw.rounded_rectangle((bar_left, y + 22, bar_right, y + 68), radius=20, fill=f"#{CARD}")
        value_right = bar_left + int((bar_right - bar_left) * float(item["value"]) / maximum)
        if item["value"] > 0:
            draw.rounded_rectangle(
                (bar_left, y + 22, max(value_right, bar_left + 25), y + 68),
                radius=20,
                fill=f"#{ACCENT}",
            )
        value_text = f"{item['value']:.1f}" if quality_score else str(int(item["value"]))
        draw.text(
            (1690, y + 27), value_text, fill=f"#{INK}", font=_font(26, bold=True), anchor="ma"
        )
    output.parent.mkdir(parents=True, exist_ok=True)
    canvas.save(output)
    return output


def _add_scorecard(document: DocxDocument, payload: ReportPayload, temp_dir: Path) -> None:
    rows, quality_score = _score_rows(payload)
    title = (
        "Sample-Limited Scorecard — 综合评分"
        if quality_score
        else "Sample-Limited Scorecard — 证据覆盖"
    )
    _add_section_band(
        document,
        "01",
        title,
        "只呈现结构化输入支持的数值；没有评分时改为证据覆盖，不推导虚构排名。",
    )
    if rows:
        chart = _create_score_chart(rows, temp_dir / "scorecard.png", quality_score=quality_score)
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        picture = paragraph.add_run().add_picture(str(chart), width=Inches(7.05))
        _set_picture_alt(
            picture,
            "Sample-limited scorecard",
            "Device bars showing batch-relative scores or approved evidence coverage.",
        )
        caption = document.add_paragraph(style="Report Caption")
        caption.add_run(
            "图 1  "
            + (
                "批次内相对评分；仅用于本批样张的结构化比较。"
                if quality_score
                else "证据覆盖数量；它衡量可发布结论的数量，不是质量排名。"
            )
        )
    else:
        _add_callout(document, "数据状态", "没有可用于评分或证据覆盖的结构化记录。")


def _add_methodology(document: DocxDocument, payload: ReportPayload) -> None:
    _add_section_band(
        document,
        "02",
        "Test Methodology — 数据、场景与边界",
        "证据优先：先按场景横向比较，再形成设备 Profile。",
    )
    table = document.add_table(rows=1, cols=2)
    _set_table_geometry(table, [3.6, 3.6])
    left, right = table.rows[0].cells
    for cell in (left, right):
        _shade_cell(cell, CARD)
        _set_cell_border(cell, color=WHITE)
    asset_count = len(payload.visual_assets)
    notes = payload.methodology_notes or [
        f"按 {len(payload.scene_results)} 个匹配场景组进行横向比较。",
        f"报告嵌入 {asset_count} 张代表样张用于人工复核。",
        "只发布通过证据等级与复核状态门槛的结论。",
    ]
    paragraph = _clear_cell(left)
    run = paragraph.add_run("数据与证据层级")
    _set_run_font(run, size=9, color=ACCENT, bold=True)
    for note in notes:
        paragraph = left.add_paragraph(style="List Bullet")
        run = paragraph.add_run(_safe_text(note))
        _set_run_font(run, size=8)
    paragraph = _clear_cell(right)
    run = paragraph.add_run("评分与机制边界")
    _set_run_font(run, size=9, color=ACCENT, bold=True)
    boundaries = payload.limitations[:4] or [
        "Rendered JPEGs do not uniquely identify internal ISP mechanisms."
    ]
    for boundary in boundaries:
        paragraph = right.add_paragraph(style="List Bullet")
        run = paragraph.add_run(_safe_text(boundary))
        _set_run_font(run, size=8)
    _add_spacer(document, 5)
    _add_callout(
        document,
        "四步判断链",
        "可见现象 → 可能原因 → 归因性质 → 用户后果。硬件提供可恢复空间，算法分配局部预算，产品偏好定义最终风格。",
        accent=True,
    )


def _resolve_asset(asset: dict[str, Any], base_dir: Path) -> Path | None:
    value = asset.get("path") or asset.get("source_path")
    if not value:
        return None
    path = Path(str(value))
    if not path.is_absolute():
        path = base_dir / path
    return path if path.is_file() else None


def _create_montage(
    assets: list[dict[str, Any]],
    base_dir: Path,
    output: Path,
) -> tuple[Path | None, int]:
    valid: list[tuple[dict[str, Any], Image.Image]] = []
    for asset in assets[:6]:
        path = _resolve_asset(asset, base_dir)
        if path is None:
            continue
        try:
            with Image.open(path) as opened_image:
                valid.append((asset, opened_image.convert("RGB")))
        except (OSError, ValueError):
            continue
    if not valid:
        return None, 0
    columns = min(3, len(valid))
    rows = math.ceil(len(valid) / columns)
    panel_width, image_height, label_height, gutter = 570, 410, 52, 14
    canvas_width = columns * panel_width + (columns - 1) * gutter
    canvas_height = rows * (image_height + label_height) + (rows - 1) * gutter
    canvas = Image.new("RGB", (canvas_width, canvas_height), f"#{CARD}")
    draw = ImageDraw.Draw(canvas)
    for index, (asset, panel_image) in enumerate(valid):
        x = (index % columns) * (panel_width + gutter)
        y = (index // columns) * (image_height + label_height + gutter)
        fitted = ImageOps.fit(
            panel_image, (panel_width, image_height), method=Image.Resampling.LANCZOS
        )
        canvas.paste(fitted, (x, y))
        draw.rectangle(
            (x, y + image_height, x + panel_width, y + image_height + label_height), fill=f"#{INK}"
        )
        label = _safe_text(
            asset.get("device_name") or asset.get("device_id") or asset.get("filename"), "Sample"
        )
        draw.text((x + 18, y + image_height + 13), label, fill="white", font=_font(23, bold=True))
    output.parent.mkdir(parents=True, exist_ok=True)
    canvas.save(output, quality=92)
    return output, len(valid)


def _scene_title(scene: dict[str, Any], index: int) -> str:
    label = _safe_text(scene.get("label") or scene.get("scene_name"))
    group_id = _safe_text(scene.get("group_id"), f"Scene {index}")
    return f"Scene {index} — {label or group_id}"


def _scene_findings(scene: dict[str, Any]) -> list[dict[str, Any]]:
    findings = scene.get("findings")
    return findings if isinstance(findings, list) else []


def _scene_observations(scene: dict[str, Any]) -> list[dict[str, Any]]:
    primary = scene.get("primary_visual") or scene.get("primary") or {}
    observations = primary.get("observations", []) if isinstance(primary, dict) else []
    if observations:
        return observations
    return _scene_findings(scene)


def _device_name(identity: str, payload: ReportPayload) -> str:
    for profile in payload.device_profiles:
        if identity in {
            _safe_text(profile.get("device_id")),
            _safe_text(profile.get("device_name")),
            _safe_text(profile.get("name")),
        }:
            return _safe_text(profile.get("device_name") or profile.get("name"), identity)
    return identity


def _add_scene_sections(
    document: DocxDocument,
    payload: ReportPayload,
    temp_dir: Path,
    base_dir: Path,
) -> None:
    assets_by_scene: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for asset in payload.visual_assets:
        assets_by_scene[_safe_text(asset.get("scene_id") or asset.get("group_id"))].append(asset)
    for index, scene in enumerate(payload.scene_results, start=1):
        scene_id = _safe_text(scene.get("group_id"), f"scene-{index}")
        observations = _scene_observations(scene)
        findings = _scene_findings(scene)
        _add_section_band(
            document,
            f"{index + 2:02d}",
            _scene_title(scene, index),
            "横向样张证据、设备观察、场景审计与效果机制解释。",
        )
        conclusion = "；".join(_safe_text(item.get("statement")) for item in findings[:4])
        audit = scene.get("audit", {}) if isinstance(scene.get("audit"), dict) else {}
        if not conclusion:
            warnings = audit.get("warnings", []) if isinstance(audit, dict) else []
            conclusion = "；".join(map(str, warnings)) or "本场景尚无通过发布门槛的结论。"
        _add_callout(document, "场景结论", conclusion, accent=True)

        montage, count = _create_montage(
            assets_by_scene.get(scene_id, []),
            base_dir,
            temp_dir / f"scene-{index:02d}.jpg",
        )
        if montage:
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            picture = paragraph.add_run().add_picture(str(montage), width=Inches(7.08))
            _set_picture_alt(
                picture,
                f"{scene_id} sample montage",
                f"Representative front-camera samples for {scene_id}; {count} labeled device panels.",
            )
            caption = document.add_paragraph(style="Report Caption")
            caption.add_run(
                f"图 {index + 1}  {scene_id} 代表样张，共 {count} 张；裁切仅用于统一版面展示。"
            )
        else:
            _add_callout(document, "样张状态", "该场景未提供可嵌入的代表样张，保留结构化观察。")

        if observations:
            table = document.add_table(rows=1, cols=3)
            _set_table_geometry(table, [1.35, 1.35, 4.5])
            headers = ("设备", "维度 / 等级", "可见观察")
            for cell, header in zip(table.rows[0].cells, headers, strict=True):
                _shade_cell(cell, INK)
                _set_cell_border(cell, color=WHITE)
                _cell_text(
                    cell,
                    header,
                    size=8,
                    color=WHITE,
                    bold=True,
                    alignment=WD_ALIGN_PARAGRAPH.CENTER,
                )
            _set_repeat_table_header(table.rows[0])
            for observation in observations[:12]:
                cells = table.add_row().cells
                _set_table_geometry(table, [1.35, 1.35, 4.5])
                identity = _safe_text(
                    observation.get("device_name") or observation.get("device_id"), "—"
                )
                dimension = _safe_text(
                    observation.get("dimension") or observation.get("grade"), "观察"
                )
                statement = _safe_text(observation.get("statement"), "未记录")
                for cell in cells:
                    _shade_cell(cell, WHITE)
                    _set_cell_border(cell)
                _cell_text(cells[0], _device_name(identity, payload), size=8, bold=True)
                _cell_text(
                    cells[1],
                    dimension,
                    size=7.5,
                    color=ACCENT,
                    bold=True,
                    alignment=WD_ALIGN_PARAGRAPH.CENTER,
                )
                _cell_text(cells[2], statement, size=8)
            _add_spacer(document, 4)

        attribution = next(
            (
                item
                for item in payload.attributions
                if scene_id in item.get("supporting_scene_ids", item.get("scene_ids", []))
            ),
            None,
        )
        if attribution:
            mechanism = _safe_text(
                attribution.get("strategy_inference")
                or attribution.get("rationale")
                or attribution.get("mechanism_hypothesis")
                or attribution.get("primary_layer")
            )
            _add_callout(
                document, "效果机制", mechanism or "归因记录存在，但没有可发布的机制表述。"
            )


def _add_attribution_and_context(
    document: DocxDocument, payload: ReportPayload, number: int
) -> int:
    _add_section_band(
        document,
        f"{number:02d}",
        "Why They Look This Way — 效果与归因",
        "区分可见现象、策略推断、机制假设与硬件上下文。",
    )
    rows = payload.attributions
    if rows:
        table = document.add_table(rows=1, cols=4)
        widths = [1.25, 2.35, 1.25, 2.35]
        _set_table_geometry(table, widths)
        for cell, header in zip(
            table.rows[0].cells, ("设备", "效果 / 推断", "归因", "替代解释"), strict=True
        ):
            _shade_cell(cell, INK)
            _set_cell_border(cell, color=WHITE)
            _cell_text(
                cell, header, size=8, color=WHITE, bold=True, alignment=WD_ALIGN_PARAGRAPH.CENTER
            )
        _set_repeat_table_header(table.rows[0])
        for item in rows:
            cells = table.add_row().cells
            _set_table_geometry(table, widths)
            alternatives = item.get("alternative_explanations", [])
            values = (
                _device_name(_safe_text(item.get("device_id"), "—"), payload),
                _safe_text(item.get("observable_effect") or item.get("rationale"), "—"),
                _safe_text(item.get("attribution") or item.get("primary_layer"), "indeterminate"),
                "；".join(map(str, alternatives)) or "—",
            )
            for cell, value in zip(cells, values, strict=True):
                _shade_cell(cell, WHITE)
                _set_cell_border(cell)
                _cell_text(cell, value, size=7.5)
    else:
        _add_callout(document, "归因状态", "当前没有通过复核的机制归因记录。")

    if payload.hardware_context:
        evidence = "；".join(
            f"{_safe_text(item.get('device_name'), 'Device')}: {_safe_text(item.get('title'), 'Source')}"
            for item in payload.hardware_context[:6]
        )
        _add_callout(document, "硬件上下文", evidence)
    if payload.external_validation:
        evidence = "；".join(
            _safe_text(item.get("title"), "Source") for item in payload.external_validation[:6]
        )
        _add_callout(document, "外部印证", evidence)
    return number + 1


def _add_failure_modes(document: DocxDocument, payload: ReportPayload, number: int) -> int:
    _add_section_band(
        document,
        f"{number:02d}",
        "Failure Modes — 失效与复核边界",
        "把风险、用户后果和下一步约束分开记录。",
    )
    entries = payload.capture_bias or []
    if entries:
        table = document.add_table(rows=1, cols=3)
        widths = [1.6, 3.2, 2.4]
        _set_table_geometry(table, widths)
        for cell, header in zip(
            table.rows[0].cells, ("状态", "用户可见风险", "建议约束"), strict=True
        ):
            _shade_cell(cell, INK)
            _set_cell_border(cell, color=WHITE)
            _cell_text(
                cell, header, size=8, color=WHITE, bold=True, alignment=WD_ALIGN_PARAGRAPH.CENTER
            )
        _set_repeat_table_header(table.rows[0])
        for item in entries[:10]:
            cells = table.add_row().cells
            _set_table_geometry(table, widths)
            values = (
                _safe_text(item.get("status"), "UNRESOLVED"),
                _safe_text(
                    item.get("assessment") or item.get("reason") or item.get("summary"),
                    "需要人工复核",
                ),
                _safe_text(
                    item.get("recommendation") or item.get("controlled_reshoot"), "增加受控重拍"
                ),
            )
            for cell, value in zip(cells, values, strict=True):
                _shade_cell(cell, WHITE)
                _set_cell_border(cell)
                _cell_text(cell, value, size=7.5)
    else:
        _add_callout(document, "复核状态", "未记录额外的拍摄偏差或失效模式。")
    _add_callout(
        document,
        "评测原则",
        "人脸可读性、光照真实性、局部自然度和风格偏好必须分开解释；缺少结构化证据时不生成单一冠军。",
        accent=True,
    )
    return number + 1


def _add_final_verdict(document: DocxDocument, payload: ReportPayload, number: int) -> int:
    _add_section_band(
        document,
        f"{number:02d}",
        "Final Verdict — 分类结论，而非默认单一冠军",
        "最终页只汇总通过证据门槛的结论、评分和限制。",
    )
    score_rows, quality_score = _score_rows(payload)
    if quality_score and score_rows:
        ordered = sorted(score_rows, key=lambda item: item["value"], reverse=True)
        table = document.add_table(rows=1, cols=3)
        _set_table_geometry(table, [3.8, 1.1, 2.3])
        for cell, header in zip(table.rows[0].cells, ("机型", "综合分", "证据状态"), strict=True):
            _shade_cell(cell, INK)
            _set_cell_border(cell, color=WHITE)
            _cell_text(
                cell, header, size=8, color=WHITE, bold=True, alignment=WD_ALIGN_PARAGRAPH.CENTER
            )
        _set_repeat_table_header(table.rows[0])
        for row in ordered:
            cells = table.add_row().cells
            _set_table_geometry(table, [3.8, 1.1, 2.3])
            values = (row["device_name"], f"{row['value']:.1f}", "批次内相对评分")
            for cell, value in zip(cells, values, strict=True):
                _shade_cell(cell, WHITE)
                _set_cell_border(cell)
                _cell_text(cell, value, size=8, bold=cell is cells[0])
    else:
        findings = _approved_findings(payload)
        _add_callout(
            document,
            "最终结论",
            "；".join(_safe_text(item.get("statement")) for item in findings[:6])
            or "当前没有足够的结构化评分或获批结论，因此不生成质量排名。",
            accent=True,
        )
    return number + 1


def _add_appendix(document: DocxDocument, payload: ReportPayload, number: int) -> None:
    _add_section_band(
        document,
        "A",
        "Appendix — 代表样张索引与限制",
        "文件索引用于复核，不把文件编号误当作跨机型统一场景编号。",
    )
    if payload.visual_assets:
        table = document.add_table(rows=1, cols=3)
        widths = [1.6, 1.8, 3.8]
        _set_table_geometry(table, widths)
        for cell, header in zip(table.rows[0].cells, ("机型", "场景", "文件名"), strict=True):
            _shade_cell(cell, INK)
            _set_cell_border(cell, color=WHITE)
            _cell_text(
                cell, header, size=8, color=WHITE, bold=True, alignment=WD_ALIGN_PARAGRAPH.CENTER
            )
        _set_repeat_table_header(table.rows[0])
        for asset in payload.visual_assets:
            cells = table.add_row().cells
            _set_table_geometry(table, widths)
            values = (
                _safe_text(asset.get("device_name") or asset.get("device_id"), "—"),
                _safe_text(asset.get("scene_id") or asset.get("group_id"), "—"),
                Path(_safe_text(asset.get("path") or asset.get("filename"), "—")).name,
            )
            for cell, value in zip(cells, values, strict=True):
                _shade_cell(cell, WHITE)
                _set_cell_border(cell)
                _cell_text(cell, value, size=7.5)
    limitations = "；".join(payload.limitations) or "没有额外限制记录。"
    _add_callout(document, "严格限制", limitations)
    _add_callout(
        document,
        "机制推断边界",
        "最终 JPEG 不能唯一识别传感器、AE、HDR 融合、局部 Tone Mapping、美颜或降噪的内部实现；机制语言应保持为证据支持的推断。",
    )


def render_docx_report(
    payload: ReportPayload,
    output: Path,
    *,
    status: str = "draft",
    template_path: Path | None = None,
) -> Path:
    """Render a reference-derived, evidence-first DOCX report."""
    template = template_path or PACKAGE_TEMPLATE
    document = Document(str(template)) if template.is_file() else Document()
    _configure_styles(document)
    section = document.sections[0]
    section.page_width = Inches(8.2681)
    section.page_height = Inches(11.6931)
    section.left_margin = Inches(0.5313)
    section.right_margin = Inches(0.5313)
    section.top_margin = Inches(0.4333)
    section.bottom_margin = Inches(0.4528)
    section.header_distance = Inches(0.1771)
    section.footer_distance = Inches(0.1771)
    for footer_table in section.footer.tables:
        _set_table_geometry(footer_table, [3.6, 3.6])

    properties = document.core_properties
    properties.title = payload.report_title or payload.project_name
    properties.subject = "Sample-based front-camera tone assessment"
    properties.author = "FroneCamera"
    properties.last_modified_by = "FroneCamera"
    properties.comments = "Generated from structured evidence; unsupported rankings are omitted."

    output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="docx-report-", dir=output.parent) as temporary:
        temp_dir = Path(temporary)
        _add_cover(document, payload, status)
        _add_executive_summary(document, payload)
        _add_scorecard(document, payload, temp_dir)
        _add_methodology(document, payload)
        _add_scene_sections(document, payload, temp_dir, output.parent)
        number = len(payload.scene_results) + 3
        number = _add_attribution_and_context(document, payload, number)
        number = _add_failure_modes(document, payload, number)
        number = _add_final_verdict(document, payload, number)
        _add_appendix(document, payload, number)
        document.save(str(output))
    return output
